"""
make_thesis_doc.py — สร้างเล่มภาคนิพนธ์ บทที่ 1–3 (+ front matter + บรรณานุกรม)
ให้โครงเหมือน "ต้นแบบ" (ภาคนิพนธ์ ม.พะเยา) แต่เนื้อหาเป็นโปรเจคของเรา
(ระบบตรวจจับตำหนิผิวเหล็กด้วยปัญญาประดิษฐ์และอธิบายได้)

    python make_thesis_doc.py
    -> docs/ภาคนิพนธ์_ระบบตรวจจับตำหนิผิวเหล็ก_บท1-3.docx
    -> figures/thesis/*.png  (ไดอาแกรมขาวดำ วาดใหม่ด้วย matplotlib)

เนื้อหาอ้างอิงจากโค้ดจริง (pipeline.py / app.py / train.py / evaluate*.py),
README.md, thesis_notes.md — ส่วน Dashboard / Severity / Alert / Roboflow API
ระบุชัดว่าเป็น "งานส่วนที่จะพัฒนาต่อ (Future Work)"
"""
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch, Polygon, Circle, Rectangle, Ellipse

from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = Path(__file__).resolve().parent
FIG = BASE / "figures" / "thesis"
FIG.mkdir(parents=True, exist_ok=True)
OUT = BASE / "docs" / "ภาคนิพนธ์_ระบบตรวจจับตำหนิผิวเหล็ก_บท1-3.docx"

for _n in ("Tahoma", "Leelawadee UI", "Angsana New"):
    try:
        font_manager.findfont(_n, fallback_to_default=False)
        plt.rcParams["font.family"] = _n
        break
    except Exception:
        continue
plt.rcParams["axes.unicode_minus"] = False

# ============================================================ ไดอาแกรมขาวดำ
GREYS = dict(box="#FFFFFF", proc="#EDEDED", alt="#E0E0E0", accent="#CFCFCF",
             soft="#F5F5F5", ec="#444444")


def _canvas(xlim, ylim, w=9.0):
    fig, ax = plt.subplots(figsize=(w, w * ylim / xlim * 0.9))
    ax.set_xlim(0, xlim); ax.set_ylim(0, ylim); ax.axis("off")
    return fig, ax


def _rbox(ax, x, y, w, h, text, fc=GREYS["box"], fs=9.5, r=True):
    style = "round,pad=0.02,rounding_size=0.35" if r else "square,pad=0.02"
    ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle=style, fc=fc, ec=GREYS["ec"], lw=1.2))
    ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=fs, zorder=5)
    return dict(cx=x + w / 2, cy=y + h / 2, l=x, r=x + w, t=y + h, b=y, w=w, h=h)


def _diam(ax, cx, cy, w, h, text, fc=GREYS["soft"], fs=8.5):
    ax.add_patch(Polygon([(cx, cy + h / 2), (cx + w / 2, cy), (cx, cy - h / 2), (cx - w / 2, cy)],
                         closed=True, fc=fc, ec=GREYS["ec"], lw=1.2))
    ax.text(cx, cy, text, ha="center", va="center", fontsize=fs, zorder=5)
    return dict(cx=cx, cy=cy, l=cx - w / 2, r=cx + w / 2, t=cy + h / 2, b=cy - h / 2)


def _term(ax, cx, cy, w, h, text, fs=9.5):
    ax.add_patch(FancyBboxPatch((cx - w / 2, cy - h / 2), w, h,
                                boxstyle=f"round,pad=0.02,rounding_size={h/2}",
                                fc=GREYS["accent"], ec=GREYS["ec"], lw=1.2))
    ax.text(cx, cy, text, ha="center", va="center", fontsize=fs, zorder=5)
    return dict(cx=cx, cy=cy, t=cy + h / 2, b=cy - h / 2, l=cx - w / 2, r=cx + w / 2)


def _ar(ax, p1, p2, txt=None, dashed=False, fs=8):
    ax.add_patch(FancyArrowPatch(p1, p2, arrowstyle="-|>", mutation_scale=13, lw=1.2,
                                 color=GREYS["ec"], ls="--" if dashed else "-",
                                 shrinkA=1, shrinkB=2))
    if txt:
        ax.text((p1[0] + p2[0]) / 2, (p1[1] + p2[1]) / 2, txt, fontsize=fs, ha="center",
                va="center", bbox=dict(fc="white", ec="none", pad=1), zorder=6)


def _save(fig, name):
    fig.savefig(FIG / name, dpi=175, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    print("  figure:", name)
    return name


# --- ภาพ: กรอบแนวคิดการศึกษา (conceptual framework) ---
def fig_framework():
    fig, ax = _canvas(100, 108, w=7.6)
    ax.text(50, 105, "กรอบแนวคิดการศึกษา", ha="center", fontsize=12, fontweight="bold")
    steps = [
        ("1. ข้อมูลนำเข้า (Input)",
         "ภาพถ่ายพื้นผิวเหล็ก (ไฟล์ภาพ หรืออัปโหลดผ่านหน้าเว็บ)"),
        ("2. การเตรียมข้อมูล (Pre-processing)",
         "แปลงเป็น BGR 3 ช่อง · แปลงเป็นภาพขาวดำเมื่อโมเดลฝึกบน grayscale\nปรับขนาดและ normalize"),
        ("3. Stage 1 — ระบุบริเวณโลหะ",
         "DMS46 → mask พื้นที่ \"Metal\" → กรอบบริเวณ\nfallback: พื้นที่โลหะ < 5% → ตรวจทั้งภาพ"),
        ("4. Stage 2 — ตรวจจับตำหนิ",
         "YOLO11 ตรวจตำหนิ 8 ชนิดในแต่ละบริเวณ\nรวมผล + Cross-region NMS + per-class threshold"),
        ("5. การอธิบายผลและการตัดสิน (Explain & Decide)",
         "กรอบตำหนิ + ป้ายชนิด (ไทย) + ค่าความมั่นใจ + ระดับความเสี่ยง + ภาพผล Stage 1\n"
         "[Future Work] Defect Area Ratio → Severity → Pass/Fail → Dashboard/Alert"),
    ]
    y = 98
    prev = None
    for head, body in steps:
        b = _rbox(ax, 4, y - 16, 92, 16, f"{head}\n{body}", GREYS["proc"], 8.6)
        if prev:
            _ar(ax, (prev["cx"], prev["b"]), (b["cx"], b["t"]))
        prev = b
        y -= 19.5
    return _save(fig, "th_framework.png")


# --- ภาพ: Flowchart ระบบหลัก ---
def fig_flow_main():
    fig, ax = _canvas(100, 150, w=6.4)
    ax.text(50, 147, "แผนผังลำดับขั้นตอนการทำงานของระบบหลัก", ha="center", fontsize=12, fontweight="bold")
    s = _term(ax, 50, 140, 34, 7, "เริ่ม")
    a = _rbox(ax, 30, 123, 40, 9, "รับภาพนำเข้า\n(cv2.imread / อัปโหลด)", GREYS["proc"], 8.5)
    d0 = _diam(ax, 50, 108, 42, 14, "เปิดไฟล์ภาพ\nสำเร็จ?")
    e0 = _rbox(ax, 80, 104, 18, 8, "แจ้ง error\nจบการทำงาน", GREYS["soft"], 8)
    b1 = _rbox(ax, 26, 88, 48, 9, "Stage 1: run_stage1()\nDMS46 → mask พื้นที่โลหะ", GREYS["proc"], 8.5)
    b2 = _rbox(ax, 24, 74, 52, 9, "build_regions(mask, image.shape)\n→ boxes, meta (metal_ratio)", GREYS["proc"], 8.5)
    d1 = _diam(ax, 50, 58, 44, 14, "metal_ratio < 0.05\nหรือไม่พบกรอบ?")
    f1 = _rbox(ax, 6, 54, 22, 8, "boxes += (0,0,W,H)\n(fallback ทั้งภาพ)", GREYS["soft"], 7.6)
    b3 = _rbox(ax, 22, 38, 56, 9, "Stage 2: run_stage2() ทุกบริเวณ\n(แปลง gray ถ้าจำเป็น · YOLO predict · กรอง conf)", GREYS["proc"], 8.2)
    b4 = _rbox(ax, 24, 25, 52, 9, "map bbox → พิกัดภาพเต็ม\ncross_region_nms(IoU 0.5)", GREYS["proc"], 8.5)
    b5 = _rbox(ax, 22, 12, 56, 9, "วาดกรอบ + ป้ายไทย + ระดับความเสี่ยง\nบันทึก *_result.jpg / *_result.json", GREYS["proc"], 8.2)
    en = _term(ax, 50, 2, 34, 6, "จบ")
    _ar(ax, (s["cx"], s["b"]), (a["cx"], a["t"]))
    _ar(ax, (a["cx"], a["b"]), (d0["cx"], d0["t"]))
    _ar(ax, (d0["r"], d0["cy"]), (e0["l"], e0["cy"]), "ไม่")
    _ar(ax, (d0["cx"], d0["b"]), (b1["cx"], b1["t"]), "ใช่")
    _ar(ax, (b1["cx"], b1["b"]), (b2["cx"], b2["t"]))
    _ar(ax, (b2["cx"], b2["b"]), (d1["cx"], d1["t"]))
    _ar(ax, (d1["l"], d1["cy"]), (f1["r"], f1["cy"]), "ใช่")
    _ar(ax, (f1["cx"], f1["b"]), (b3["l"] + 6, b3["t"]))
    _ar(ax, (d1["cx"], d1["b"]), (b3["cx"], b3["t"]), "ไม่")
    _ar(ax, (b3["cx"], b3["b"]), (b4["cx"], b4["t"]))
    _ar(ax, (b4["cx"], b4["b"]), (b5["cx"], b5["t"]))
    _ar(ax, (b5["cx"], b5["b"]), (en["cx"], en["t"]))
    return _save(fig, "th_flow_main.png")


# --- ภาพ: Flowchart การเตรียมภาพ (pre-processing) ---
def fig_flow_pre():
    fig, ax = _canvas(100, 96, w=6.2)
    ax.text(50, 93, "แผนผังลำดับขั้นตอนการปรับแต่งภาพ (Pre-processing)", ha="center", fontsize=11.5, fontweight="bold")
    s = _term(ax, 50, 86, 30, 6, "เริ่ม")
    steps = [
        "รับ array ภาพจากผู้ใช้ (อาจเป็น grayscale / RGBA / RGB)",
        "_to_bgr(): แปลงให้เป็น BGR 3 ช่องเสมอ",
        "ตรวจ checkpoint ของ Stage 2 ว่าเทรนบน grayscale หรือไม่",
        "ถ้าใช่ → แปลง crop เป็นภาพขาวดำ (กัน train/serve skew)",
        "Stage 1: resize ด้านยาว = 512 (รักษาสัดส่วน) + ImageNet normalize",
        "ส่งภาพที่เตรียมแล้วเข้าสู่ Stage 1 / Stage 2",
    ]
    y = 78
    prev = s
    for t in steps:
        b = _rbox(ax, 12, y - 8, 76, 8, t, GREYS["proc"], 8.3)
        _ar(ax, (prev["cx"], prev["b"]), (b["cx"], b["t"]))
        prev = b; y -= 12
    en = _term(ax, 50, y, 30, 6, "จบ")
    _ar(ax, (prev["cx"], prev["b"]), (en["cx"], en["t"]))
    return _save(fig, "th_flow_pre.png")


# --- ภาพ: dataset class distribution ---
def fig_class_dist():
    classes = ["crazing", "inclusion", "patches", "pitted_surface",
               "rolled-in_scale", "scratches", "rust", "crack"]
    counts = [555, 997, 902, 714, 504, 1000, 705, 1125]   # จาก README (instance ต่อคลาส, โดยประมาณ)
    fig, ax = plt.subplots(figsize=(8, 3.6))
    bars = ax.bar(range(len(classes)), counts, color="#B0B0B0", edgecolor="#444", linewidth=1)
    ax.set_xticks(range(len(classes)))
    ax.set_xticklabels(classes, rotation=25, ha="right", fontsize=9)
    ax.set_ylabel("จำนวน instance (โดยประมาณ)", fontsize=9)
    ax.set_title("การกระจายของคลาสตำหนิใน merged_dataset (8 คลาส)", fontsize=11, fontweight="bold")
    for b, c in zip(bars, counts):
        ax.text(b.get_x() + b.get_width() / 2, c + 15, str(c), ha="center", fontsize=8)
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    return _save(fig, "th_class_dist.png")


def gen_reused_figs():
    """ยืมฟังก์ชันวาดไดอาแกรมจากสคริปต์เดิม แล้วบังคับให้เป็นโทนขาวดำ"""
    names = {}
    grey = {"a": "#ECECEC", "s1": "#DCDCDC", "s2": "#F2F2F2", "data": "#E4E4E4",
            "out": "#F6F6F6", "note": "#FAFAFA", "ext": "#EFEFEF", "sys": "#D0D0D0"}
    try:
        import make_uml_doc as U
        U.CL.update(grey)
        U.DIA = FIG                     # เขียนรูปลงโฟลเดอร์เดียวกัน
        names["usecase"] = U.d1()
        names["context"] = U.d2()
        names["c4l1"] = U.d3()
        names["c4l2"] = U.d4()
        names["erd"] = U.d5()
        names["class"] = U.d6()
        names["sequence"] = U.d7()
        names["activity"] = U.d8()
        names["state"] = U.d9()
        names["deploy"] = U.d10()
    except Exception as e:
        print("  (ข้าม make_uml_doc:", e, ")")
    try:
        import make_diagrams_doc as G
        G.C.update(grey)
        G.DIA = FIG
        names["arch"] = G.d1()
        names["stage1"] = G.d2()
        names["stage2"] = G.d3()
        names["dataprep"] = G.d4()
    except Exception as e:
        print("  (ข้าม make_diagrams_doc:", e, ")")
    try:
        import make_design_doc as D
        D.DIA = FIG
        D.make_wireframe()
        names["wireframe"] = "ui_wireframe.png"
    except Exception as e:
        print("  (ข้าม make_design_doc:", e, ")")
    return names


# ============================================================ docx helpers
TH_FONT = "TH Sarabun New"
BODY_PT, HEAD_PT = 16, 18


def _set_cell_font(cell, size=14, bold=False):
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.name = TH_FONT
            r.font.size = Pt(size)
            r.font.bold = bold
            rpr = r._element.get_or_add_rPr()
            rf = OxmlElement("w:rFonts")
            rf.set(qn("w:cs"), TH_FONT)
            rpr.append(rf)


class Doc:
    def __init__(self):
        self.d = Document()
        self.fig_n = 0
        self.tab_n = 0
        self.figs = []        # (number, caption)
        self.tabs = []
        st = self.d.styles["Normal"]
        st.font.name = TH_FONT
        st.font.size = Pt(BODY_PT)
        st.element.rPr.rFonts.set(qn("w:cs"), TH_FONT)
        for i in range(1, 5):
            hs = self.d.styles[f"Heading {i}"]
            hs.font.name = TH_FONT
            hs.font.size = Pt(HEAD_PT if i == 1 else max(BODY_PT, HEAD_PT - i))
            hs.font.color.rgb = RGBColor(0, 0, 0)
            hs.font.bold = True
            hs.element.rPr.rFonts.set(qn("w:cs"), TH_FONT)
        for s in self.d.sections:
            s.left_margin = Cm(3.0); s.right_margin = Cm(2.0)
            s.top_margin = Cm(2.5); s.bottom_margin = Cm(2.5)

    # ---- content ----
    def h(self, text, level=1, num=True):
        p = self.d.add_heading(text, level=level)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        return p

    def chapter(self, no, title):
        self.d.add_page_break()
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"บทที่ {no}")
        r.bold = True; r.font.size = Pt(HEAD_PT); r.font.name = TH_FONT
        p2 = self.d.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(title)
        r2.bold = True; r2.font.size = Pt(HEAD_PT); r2.font.name = TH_FONT
        p2.paragraph_format.space_after = Pt(12)
        # heading ที่ไม่แสดงบนหน้า แต่เข้า TOC
        hd = self.d.add_heading(f"บทที่ {no} {title}", level=1)
        hd.runs[0].font.size = Pt(1)
        hd.runs[0].font.color.rgb = RGBColor(255, 255, 255)

    def p(self, text, indent=True, align=None):
        par = self.d.add_paragraph(text)
        par.paragraph_format.first_line_indent = Cm(1.27 if indent else 0)
        par.paragraph_format.space_after = Pt(6)
        par.paragraph_format.line_spacing = 1.15
        if align == "c":
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return par

    def bullets(self, items, style="List Bullet"):
        for it in items:
            par = self.d.add_paragraph(it, style=style)
            par.paragraph_format.space_after = Pt(2)
            for r in par.runs:
                r.font.name = TH_FONT; r.font.size = Pt(BODY_PT)

    def numbered(self, items):
        self.bullets(items, style="List Number")

    def _seq_caption(self, label, text, bold=False, align="c", space_after=12, space_before=0):
        cap = self.d.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER if align == "c" else WD_ALIGN_PARAGRAPH.LEFT
        cap.paragraph_format.space_after = Pt(space_after)
        cap.paragraph_format.space_before = Pt(space_before)
        r0 = cap.add_run(label + " "); r0.font.name = TH_FONT; r0.font.size = Pt(14); r0.bold = bold
        run = cap.add_run()
        b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
        it.text = f' SEQ {label} \\* ARABIC '
        sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
        t = OxmlElement("w:t"); t.text = "0"
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
        for el in (b, it, sep, t, e):
            run._r.append(el)
        run.font.name = TH_FONT; run.font.size = Pt(14); run.bold = bold
        r2 = cap.add_run("  " + text); r2.font.name = TH_FONT; r2.font.size = Pt(14); r2.bold = bold
        return cap

    def figure(self, img, caption, width=6.2):
        self.fig_n += 1
        par = self.d.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before = Pt(6)
        par.add_run().add_picture(str(FIG / img), width=Inches(width))
        self._seq_caption("ภาพ", caption)
        self.figs.append((self.fig_n, caption))

    def table(self, headers, rows, caption, col_w=None):
        self.tab_n += 1
        self._seq_caption("ตาราง", caption, bold=True, space_before=8, space_after=2)
        t = self.d.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        for i, hcell in enumerate(headers):
            t.rows[0].cells[i].text = str(hcell)
            _set_cell_font(t.rows[0].cells[i], 13, bold=True)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = str(v)
                _set_cell_font(cells[i], 13)
        if col_w:
            for i, w in enumerate(col_w):
                for row in t.rows:
                    row.cells[i].width = Cm(w)
        self.d.add_paragraph().paragraph_format.space_after = Pt(6)
        self.tabs.append((self.tab_n, caption))
        return t

    def toc_field(self, instr, hint):
        p = self.d.add_paragraph()
        run = p.add_run()
        b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
        sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
        run._r.append(b); run._r.append(it); run._r.append(sep)
        hr = p.add_run(hint); hr.italic = True; hr.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
        run2 = p.add_run(); run2._r.append(e)

    def page_numbers(self):
        for s in self.d.sections:
            fp = s.footer.paragraphs[0]
            fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = fp.add_run()
            b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = "PAGE"
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
            run._r.append(b); run._r.append(it); run._r.append(e)
            run.font.size = Pt(13); run.font.name = TH_FONT

    def save(self):
        OUT.parent.mkdir(parents=True, exist_ok=True)
        self.d.save(str(OUT))
        print("\nเขียนเล่ม:", OUT.relative_to(BASE))


# ============================================================ front matter + บท 1-2
def _center_line(d, txt, sz=16, bold=False):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(txt); r.bold = bold; r.font.size = Pt(sz); r.font.name = TH_FONT
    return p


def build_front_and_ch12(figs):
    D = Doc()
    d = D.d

    # ---- ปกนอก ----
    for _ in range(4):
        d.add_paragraph()
    _center_line(d, "ระบบตรวจจับตำหนิผิวเหล็กด้วยปัญญาประดิษฐ์และอธิบายได้", 20, True)
    _center_line(d, "Explainable Steel Surface Defect Detection", 16)
    for _ in range(6):
        d.add_paragraph()
    _center_line(d, "นายวีรกร วงศ์ษารัฐ  รหัสนิสิต 67021253", 16)
    _center_line(d, "นายกฤตเมธ ป้องตัน  รหัสนิสิต 67020746", 16)
    for _ in range(6):
        d.add_paragraph()
    for t in ["โครงงานนี้เป็นส่วนหนึ่งของรายวิชาการเตรียมพร้อมสำหรับโครงงาน (225291)",
              "สาขาวิชาวิทยาการคอมพิวเตอร์  คณะเทคโนโลยีสารสนเทศและการสื่อสาร",
              "ภาคเรียนที่ 2  ปีการศึกษา 2569  มหาวิทยาลัยพะเยา"]:
        _center_line(d, t, 15)

    # ---- ปกใน / หน้าอนุมัติ ----
    d.add_page_break()
    for _ in range(3):
        d.add_paragraph()
    _center_line(d, "ระบบตรวจจับตำหนิผิวเหล็กด้วยปัญญาประดิษฐ์และอธิบายได้", 19, True)
    _center_line(d, "Explainable Steel Surface Defect Detection", 15)
    for _ in range(4):
        d.add_paragraph()
    _center_line(d, "นายวีรกร วงศ์ษารัฐ  รหัสนิสิต 67021253", 15)
    _center_line(d, "นายกฤตเมธ ป้องตัน  รหัสนิสิต 67020746", 15)
    for _ in range(6):
        d.add_paragraph()
    for t in ["อาจารย์ที่ปรึกษาโครงงาน", "",
              ".............................................................",
              "(อาจารย์ ธนวัฒน์ แซ่เอียบ)",
              "วันที่ ......... เดือน ................................ พ.ศ. ............"]:
        _center_line(d, t, 15)

    # ---- บทคัดย่อ ----
    d.add_page_break()
    hb = d.add_heading("บทคัดย่อ", level=1); hb.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for k, v in [
        ("ชื่อเรื่อง", "ระบบตรวจจับตำหนิผิวเหล็กด้วยปัญญาประดิษฐ์และอธิบายได้"),
        ("ผู้จัดทำ", "นายวีรกร วงศ์ษารัฐ  และ  นายกฤตเมธ ป้องตัน"),
        ("อาจารย์ที่ปรึกษา", "อาจารย์ ธนวัฒน์ แซ่เอียบ"),
        ("สาขาวิชา", "วิทยาการคอมพิวเตอร์  มหาวิทยาลัยพะเยา  ปีการศึกษา 2569"),
        ("คำสำคัญ", "การตรวจจับตำหนิพื้นผิวเหล็ก, การตรวจจับวัตถุ, YOLO, material segmentation "
                    "(DMS46), ปัญญาประดิษฐ์ที่อธิบายได้ (XAI), ระบบสองขั้น")]:
        par = d.add_paragraph()
        r1 = par.add_run(f"{k}  :  "); r1.bold = True; r1.font.name = TH_FONT; r1.font.size = Pt(BODY_PT)
        r2 = par.add_run(v); r2.font.name = TH_FONT; r2.font.size = Pt(BODY_PT)
    for t in [
        "การตรวจสอบคุณภาพพื้นผิวเหล็กในอุตสาหกรรมส่วนใหญ่ยังพึ่งพาการตรวจด้วยสายตาของพนักงาน ซึ่งมี "
        "ข้อจำกัดด้านความล้าและความไม่สม่ำเสมอของเกณฑ์ ขณะที่ระบบตรวจจับด้วยการเรียนรู้เชิงลึกที่มีอยู่ "
        "มักทำงานแบบ \"กล่องดำ\" ที่บอกเพียงผลว่าผ่านหรือไม่ผ่านโดยไม่อธิบายเหตุผล โครงงานนี้จึงพัฒนา "
        "ระบบตรวจจับตำหนิผิวเหล็กแบบสองขั้น (2-Stage) ที่ให้ผลลัพธ์อธิบายได้เชิงตำแหน่งและเชิงปริมาณ",
        "ระบบประกอบด้วย ขั้นที่ 1 การระบุบริเวณที่เป็นวัสดุโลหะด้วยแบบจำลอง DMS46 (Apple Dense Material "
        "Segmentation) พร้อมกลไก fallback ตรวจทั้งภาพเมื่อไม่พบบริเวณโลหะ และ ขั้นที่ 2 การตรวจจับและ "
        "จำแนกตำหนิ 8 ประเภท (crazing, inclusion, patches, pitted surface, rolled-in scale, scratches, "
        "rust, crack) ด้วยแบบจำลอง YOLO11 ที่ฝึกเองบนชุดข้อมูลรวมจาก NEU Surface Defect Database ชุด "
        "ข้อมูลสนิม และชุดข้อมูลรอยแตก โดยฝึกบนภาพระดับเทาเพื่อลดการเรียนรู้ทางลัดจากโทนสี ผลลัพธ์แสดง "
        "ผ่านหน้าเว็บเป็นกรอบตำหนิ ป้ายชนิดภาษาไทย ค่าความมั่นใจ และระดับความเสี่ยงของแต่ละคลาส",
        "การประเมินบนชุดทดสอบพบว่าแบบจำลองขั้นที่ 2 ได้ค่า mAP@0.5 เฉลี่ย 0.867 ± 0.010 (ทดลองซ้ำ 4 "
        "รอบ), mAP@0.5–0.95 เท่ากับ 0.536 และ recall 0.809 โดยคลาส crack ยังมีประสิทธิภาพต่ำที่สุด "
        "การทดสอบกับภาพถ่ายจริงระดับสถานที่แสดงช่องว่างเชิงโดเมนที่ชัดเจน และการวัดขั้นที่ 1 เชิงตัวเลข "
        "ชี้ว่าการใช้ material segmentation เป็น front-end ยังไม่คุ้มค่า จึงรายงานเป็น ablation เชิงลบ "
        "ส่วนการประเมินระดับความรุนแรง การตัดสิน Pass/Fail แดชบอร์ด และการแจ้งเตือน เป็นขอบเขตงานพัฒนาต่อ",
    ]:
        D.p(t)

    # ---- กิตติกรรมประกาศ ----
    d.add_page_break()
    ha = d.add_heading("กิตติกรรมประกาศ", level=1); ha.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for t in [
        "โครงงานเรื่อง ระบบตรวจจับตำหนิผิวเหล็กด้วยปัญญาประดิษฐ์และอธิบายได้ สำเร็จลุล่วงได้ด้วยความกรุณา "
        "และความช่วยเหลือจากหลายฝ่าย คณะผู้จัดทำขอขอบพระคุณ อาจารย์ ธนวัฒน์ แซ่เอียบ อาจารย์ที่ปรึกษา "
        "โครงงาน ที่ให้คำปรึกษา คำแนะนำ และข้อเสนอแนะตั้งแต่การกำหนดแนวทาง การออกแบบและพัฒนาระบบ ตลอดจน "
        "การตรวจสอบและแก้ไขข้อบกพร่อง",
        "ขอขอบพระคุณคณาจารย์และบุคลากรสาขาวิชาวิทยาการคอมพิวเตอร์ที่ประสิทธิ์ประสาทความรู้ และขอบคุณ "
        "ครอบครัวและเพื่อน ๆ ที่ให้การสนับสนุนและกำลังใจตลอดการจัดทำโครงงาน",
    ]:
        D.p(t)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("คณะผู้จัดทำ"); r.font.name = TH_FONT; r.font.size = Pt(BODY_PT)

    # ---- สารบัญ ----
    d.add_page_break()
    _center_line(d, "สารบัญ", 18, True)
    D.toc_field('TOC \\o "1-3" \\h \\z \\u', "  [คลิกขวา > Update Field หรือกด F9 เพื่อสร้างสารบัญ]")
    d.add_page_break()
    _center_line(d, "สารบัญตาราง", 18, True)
    D.toc_field('TOC \\h \\z \\c "ตาราง"', "  [กด F9 เพื่อสร้างสารบัญตาราง]")
    d.add_page_break()
    _center_line(d, "สารบัญรูปภาพ", 18, True)
    D.toc_field('TOC \\h \\z \\c "ภาพ"', "  [กด F9 เพื่อสร้างสารบัญรูปภาพ]")

    # ============================== บทที่ 1 ==============================
    D.chapter(1, "บทนำ")
    D.p("ในอุตสาหกรรมแปรรูปและผลิตเหล็ก คุณภาพของพื้นผิวเหล็กเป็นปัจจัยชี้ขาดที่ส่งผลต่อความแข็งแรงเชิง "
        "วิศวกรรม ความปลอดภัย และมาตรฐานผลิตภัณฑ์ ตำหนิบนผิวเหล็ก เช่น รอยขีดข่วน รอยแตกร้าว ผิวเป็นหลุม "
        "สะเก็ดฝังจากการรีด หรือสิ่งแปลกปลอมบนผิว มักเกิดจากความผิดปกติในกระบวนการรีดร้อนหรือรีดเย็น หาก "
        "ตรวจไม่พบตั้งแต่ต้นทางจะหลุดรอดไปถึงผู้ใช้งานปลายทางได้")
    D.p("โครงงานนี้พัฒนา “ระบบตรวจจับตำหนิผิวเหล็กด้วยปัญญาประดิษฐ์และอธิบายได้” โดยบูรณาการแบบจำลอง "
        "ตรวจจับวัตถุเข้ากับแนวคิดปัญญาประดิษฐ์ที่อธิบายได้ (Explainable AI) เพื่อให้ระบบไม่เพียงตรวจจับ "
        "ตำแหน่งและชนิดของตำหนิ แต่ยังอธิบายผลการตัดสินใจเชิงตำแหน่งและปริมาณได้")

    D.h("1. ความเป็นมาและความสำคัญของปัญหา", 2)
    for t in [
        "จากการสังเกตในชีวิตประจำวัน คณะผู้จัดทำพบว่าผลิตภัณฑ์ที่ทำจากเหล็กซึ่งวางจำหน่ายทั่วไป เช่น "
        "เหล็กเส้นก่อสร้างและแผงเหล็กในบ้าน บางชิ้นมีตำหนิ เช่น รอยขีดข่วนหรือรอยแตกร้าว หลุดรอดออกมาถึง "
        "มือผู้บริโภค ซึ่งควรถูกคัดแยกออกตั้งแต่ขั้นตอนการตรวจสอบคุณภาพก่อนส่งมอบ",
        "กระบวนการตรวจสอบคุณภาพในหลายโรงงานยังพึ่งพาการตรวจด้วยสายตาของพนักงาน (Visual Inspection) เป็น "
        "หลัก ซึ่งมีข้อจำกัดด้านความล่าช้าและความคลาดเคลื่อนจากความเหนื่อยล้าและดุลยพินิจของผู้ตรวจแต่ละคน "
        "จึงเป็นช่องว่างที่ทำให้ตำหนิบางส่วนตรวจไม่พบ",
        "แม้จะมีการนำระบบคอมพิวเตอร์วิทัศน์และการเรียนรู้เชิงลึกมาช่วยตรวจจับตำหนิ แต่แบบจำลองส่วนใหญ่ยัง "
        "เป็น “กล่องดำ” (Black-box) ที่บอกเพียงผลว่า “ผ่าน” หรือ “ไม่ผ่าน” โดยไม่อธิบายเหตุผล ทำให้ผู้ "
        "ควบคุมคุณภาพขาดความมั่นใจและยังต้องตรวจซ้ำด้วยคน ซึ่งขัดกับเป้าหมายที่ต้องการลดภาระงาน",
        "จากปัญหาทั้งสองส่วน คณะผู้จัดทำจึงพัฒนาระบบที่นำคอมพิวเตอร์วิทัศน์มาตรวจจับตำหนิบนผิวเหล็กโดย "
        "อัตโนมัติ ร่วมกับการอธิบายผลเชิงตำแหน่ง (Bounding Box) ค่าความมั่นใจ ระดับความเสี่ยงของแต่ละคลาส "
        "และภาพแสดงบริเวณที่ระบบพิจารณา เพื่อเพิ่มความน่าเชื่อถือของผลการตรวจสอบ",
        "ในช่วงที่ผ่านมา คณะผู้จัดทำได้พัฒนาต้นแบบ (Prototype) ที่รับภาพชิ้นงานเหล็กและประมวลผลตรวจจับ "
        "ตำหนิได้จริง โดยฝึกและทดสอบกับชุดข้อมูลมาตรฐาน NEU Surface Defect Database ร่วมกับชุดข้อมูลสนิม "
        "และรอยแตกที่จัดการผ่านแพลตฟอร์ม Roboflow ซึ่งเป็นจุดตั้งต้นสำหรับการพัฒนาส่วนอธิบายผลและการ "
        "ประเมินระดับความรุนแรงในขั้นต่อไป",
    ]:
        D.p(t)

    D.h("2. วัตถุประสงค์ของการวิจัย", 2)
    D.numbered([
        "เพื่อออกแบบและพัฒนาระบบตรวจจับตำหนิบนพื้นผิวเหล็กด้วยปัญญาประดิษฐ์ในรูปแบบระบบสองขั้น "
        "(material localization + defect detection)",
        "เพื่อประยุกต์แนวคิดปัญญาประดิษฐ์ที่อธิบายได้ (XAI) ในการอธิบายผลการตรวจจับเชิงตำแหน่งและเชิง "
        "ปริมาณ และวางกรอบการประเมินระดับความรุนแรง (Severity Assessment) ของตำหนิ",
        "เพื่อประเมินประสิทธิภาพของแบบจำลองด้วยค่า Precision, Recall, mean Average Precision (mAP) และ "
        "เวลาในการประมวลผลต่อภาพ พร้อมประเมินความพึงพอใจของผู้ใช้งานที่มีต่อระบบต้นแบบ",
    ])

    D.h("3. แนวคิดและหลักการ", 2)
    D.p("โครงงานนี้พัฒนาระบบตรวจจับและอธิบายผลตำหนิบนผิวเหล็กจากภาพถ่าย โดยใช้เทคนิคการประมวลผลภาพร่วม "
        "กับการเรียนรู้เชิงลึก กรอบแนวคิดของระบบแบ่งเป็น 5 ขั้นตอนหลัก ดังภาพ 1")
    D.figure(figs["framework"], "แสดงกรอบแนวคิดการศึกษา", width=5.2)
    D.bullets([
        "ข้อมูลนำเข้า (Input): ภาพถ่ายพื้นผิวเหล็กจากไฟล์ภาพหรือการอัปโหลดผ่านหน้าเว็บ",
        "การเตรียมข้อมูล (Pre-processing): แปลงภาพเป็น BGR สามช่อง ปรับขนาดและ normalize สำหรับขั้นที่ 1 "
        "และแปลงเป็นภาพระดับเทาก่อนเข้าขั้นที่ 2 เมื่อแบบจำลองฝึกบน grayscale",
        "Stage 1 — การระบุบริเวณโลหะ: ใช้ DMS46 หาพื้นที่วัสดุ “Metal” แล้วแปลงเป็นกรอบบริเวณ หากพบพื้นที่ "
        "โลหะน้อยกว่าเกณฑ์จะใช้กลไก fallback ตรวจทั้งภาพ",
        "Stage 2 — การตรวจจับตำหนิ: ใช้ YOLO11 ตรวจจับและจำแนกตำหนิ 8 ประเภทในแต่ละบริเวณ รวมผลข้ามบริเวณ "
        "ด้วย Cross-region NMS และกรองด้วย per-class confidence threshold",
        "การอธิบายผลและการตัดสิน: แสดงกรอบตำหนิ ป้ายชนิดภาษาไทย ค่าความมั่นใจ ระดับความเสี่ยง และภาพผล "
        "ของขั้นที่ 1 ส่วนการคำนวณ Defect Area Ratio เพื่อจัดระดับ Severity และการตัดสิน Pass/Fail พร้อม "
        "แดชบอร์ดและการแจ้งเตือน เป็นขอบเขตงานพัฒนาต่อ",
    ])

    D.h("4. ขอบเขตของงานวิจัย", 2)
    D.bullets([
        "ขอบเขตด้านข้อมูล: ตำหนิผิวเหล็ก 8 ประเภท — 6 ประเภทจาก NEU Surface Defect Database "
        "และเพิ่ม rust กับ crack จากชุดข้อมูลภายนอกที่จัดการผ่าน Roboflow",
        "ขอบเขตด้านลักษณะการตรวจสอบ: ตรวจเฉพาะตำหนิภายนอกบนผิวหน้าของชิ้นงาน ไม่ครอบคลุมโครงสร้างระดับ "
        "ไมโครภายในเนื้อเหล็ก",
        "ขอบเขตด้านฮาร์ดแวร์และซอฟต์แวร์: พัฒนาและประมวลผลบนคอมพิวเตอร์เครื่องเดียว (Windows 11 + NVIDIA "
        "RTX 3050 6GB) ไม่ต่อพ่วงเครื่องจักรอุตสาหกรรมจริงและไม่ทำงานแบบเรียลไทม์บนสายพาน",
        "ขอบเขตด้าน Explainability: อธิบายผลเชิงปริมาณและเชิงภาพผ่าน Bounding Box ค่าความมั่นใจ ระดับ "
        "ความเสี่ยงรายคลาส และภาพแสดงบริเวณของขั้นที่ 1",
        "ขอบเขตด้านการประเมิน: วัดประสิทธิภาพด้วย Precision, Recall, mAP@0.5, mAP@0.5–0.95 และเวลาในการ "
        "ประมวลผลต่อภาพ ร่วมกับแบบประเมินความพึงพอใจของผู้ใช้งาน",
        "งานส่วนที่กำหนดเป็น Future Work: การคำนวณ Severity (Low/Medium/High), การตัดสิน Pass/Fail, การ "
        "บันทึกลงฐานข้อมูล, แดชบอร์ดสรุปสถิติ และการแจ้งเตือนเมื่อพบตำหนิร้ายแรง",
    ])

    D.h("5. ประโยชน์ที่คาดว่าจะได้รับ", 2)
    D.bullets([
        "เชิงปฏิบัติ: ลดภาระงานตรวจสอบด้วยสายตา เพิ่มความสม่ำเสมอของเกณฑ์ และลดโอกาสที่ตำหนิจะหลุดรอด",
        "เชิงวิชาการ: เป็นกรณีศึกษาการประยุกต์ระบบสองขั้นและแนวคิด Explainable AI ในงานตรวจสอบคุณภาพ "
        "พื้นผิวอุตสาหกรรม รวมถึงบทเรียนเชิงลบเรื่องช่องว่างเชิงโดเมนของชุดข้อมูลแบบแล็บ",
        "เชิงต่อยอด: โครงสร้างระบบและเอกสารออกแบบสามารถนำไปพัฒนาเป็นระบบที่มีฐานข้อมูล แดชบอร์ด และการ "
        "แจ้งเตือนสำหรับใช้งานจริงในอนาคต",
    ])

    D.h("6. คำศัพท์เฉพาะและคำจำกัดความ", 2)
    for term, defn in [
        ("ปัญญาประดิษฐ์ (AI)", "เทคโนโลยีที่ทำให้คอมพิวเตอร์เรียนรู้ วิเคราะห์ และตัดสินใจจากข้อมูล"),
        ("คอมพิวเตอร์วิทัศน์ (Computer Vision)", "ศาสตร์การประมวลผลข้อมูลภาพเพื่อให้คอมพิวเตอร์เข้าใจภาพ"),
        ("การตรวจจับวัตถุ (Object Detection)", "การระบุตำแหน่งและประเภทของวัตถุในภาพ พร้อมสร้างกรอบล้อมรอบ"),
        ("YOLO (You Only Look Once)", "สถาปัตยกรรมตรวจจับวัตถุแบบขั้นตอนเดียวที่ประมวลผลรวดเร็ว ใช้ YOLO11"),
        ("Material Segmentation / DMS46", "การแบ่งส่วนภาพตามชนิดวัสดุ ใช้ DMS46 ระบุบริเวณวัสดุโลหะ"),
        ("Bounding Box", "กรอบสี่เหลี่ยมที่ระบุตำแหน่งและขอบเขตของตำหนิที่ตรวจพบ"),
        ("Confidence Score", "ค่าความมั่นใจของแบบจำลองต่อผลการตรวจจับแต่ละรายการ"),
        ("Explainable AI (XAI)", "แนวทางทำให้ผลลัพธ์ของระบบ AI อธิบายและตีความได้"),
        ("Defect Area Ratio", "อัตราส่วนพื้นที่ตำหนิเทียบกับพื้นที่บริเวณที่ตรวจสอบ (ใช้ประเมิน Severity — Future Work)"),
        ("Severity Assessment", "การประเมินระดับความรุนแรงของตำหนิจากขนาดและชนิด (Future Work)"),
        ("mAP (mean Average Precision)", "ค่าเฉลี่ยความแม่นยำของการตรวจจับที่ระดับ IoU ที่กำหนด"),
        ("NMS (Non-Maximum Suppression)", "การตัดกรอบตรวจจับที่ซ้ำซ้อน เก็บเฉพาะกรอบที่มั่นใจสูงสุด"),
        ("Roboflow", "แพลตฟอร์มจัดการชุดข้อมูลภาพ การติดป้ายกำกับ และการทำ Data Augmentation"),
    ]:
        par = d.add_paragraph()
        r1 = par.add_run(term + "  "); r1.bold = True; r1.font.name = TH_FONT; r1.font.size = Pt(BODY_PT)
        r2 = par.add_run("— " + defn); r2.font.name = TH_FONT; r2.font.size = Pt(BODY_PT)
        par.paragraph_format.space_after = Pt(4)

    # ============================== บทที่ 2 ==============================
    D.chapter(2, "เอกสารและงานวิจัยที่เกี่ยวข้อง")
    D.p("บทนี้ทบทวนทฤษฎี หลักการ เครื่องมือ และงานวิจัยที่เกี่ยวข้องกับการตรวจจับตำหนิบนพื้นผิวเหล็กด้วย "
        "ปัญญาประดิษฐ์ เพื่อใช้เป็นพื้นฐานในการออกแบบระบบในบทที่ 3")

    D.h("1. การตรวจสอบคุณภาพพื้นผิวเหล็กและผลกระทบของตำหนิ", 2)
    D.p("คุณภาพพื้นผิวเหล็กมีผลต่อความแข็งแรงและอายุการใช้งานของชิ้นงาน ตำหนิแต่ละชนิดมีระดับความเสี่ยง "
        "ต่างกัน โครงงานนี้จัดระดับความเสี่ยงของตำหนิ 8 ประเภทไว้ดังตาราง 1")
    D.table(["id", "คลาส (อังกฤษ)", "ชื่อไทย", "แหล่งข้อมูล", "ระดับความเสี่ยง"], [
        [0, "crazing", "รอยแตกลายงา", "NEU", "ปานกลาง–สูง"],
        [1, "inclusion", "สิ่งแปลกปลอมฝังใน", "NEU", "ปานกลาง"],
        [2, "patches", "รอยแผ่น/ผิวลอก", "NEU", "ต่ำ–ปานกลาง"],
        [3, "pitted_surface", "ผิวขรุขระเป็นหลุม", "NEU", "ปานกลาง"],
        [4, "rolled-in_scale", "สะเก็ดฝังจากการรีด", "NEU", "ปานกลาง"],
        [5, "scratches", "รอยขีดข่วน", "NEU", "ต่ำ"],
        [6, "rust", "สนิม", "Rust dataset", "สูง"],
        [7, "crack", "รอยแตกร้าว", "Crack dataset", "สูง"],
    ], "คลาสตำหนิ 8 ประเภทและระดับความเสี่ยงที่ใช้ในระบบ", col_w=[1.2, 3.4, 3.6, 3.2, 3.0])

    D.h("2. ปัญญาประดิษฐ์และการเรียนรู้ของเครื่อง (AI & Machine Learning)", 2)
    D.bullets([
        "การเรียนรู้ของเครื่อง (Machine Learning): การสอนให้ระบบเรียนรู้รูปแบบจากชุดข้อมูลตัวอย่างเพื่อ "
        "สร้างแบบจำลองการตัดสินใจ",
        "การเรียนรู้เชิงลึก (Deep Learning): เทคนิคที่ใช้โครงข่ายประสาทเทียมหลายชั้นประมวลผลข้อมูลภาพที่ "
        "ซับซ้อน ทำให้ระบบเรียนรู้คุณลักษณะของตำหนิได้เองโดยไม่ต้องกำหนดกฎด้วยมือ",
    ])

    D.h("3. คอมพิวเตอร์วิทัศน์สำหรับการตรวจสอบคุณภาพ", 2)
    D.bullets([
        "การจำแนกภาพ (Image Classification): จัดกลุ่มภาพทั้งภาพว่าผ่านเกณฑ์ (OK) หรือมีตำหนิ (NG)",
        "การตรวจจับวัตถุ (Object Detection): ระบุทั้งประเภทและตำแหน่งของตำหนิด้วย Bounding Box โครงงาน "
        "เลือกใช้การตรวจจับวัตถุเป็นแกนหลัก",
    ])

    D.h("4. โครงข่ายประสาทเทียมแบบคอนโวลูชัน (CNN)", 2)
    D.bullets([
        "ชั้นคอนโวลูชัน (Convolutional Layer): สกัดคุณลักษณะเด่น เช่น เส้นขอบ ความขรุขระ หรือรอยแตก",
        "ชั้นพูลลิง (Pooling Layer): ลดขนาดข้อมูลเพื่อเพิ่มความเร็ว โดยคงคุณลักษณะสำคัญไว้",
        "หัวตรวจจับ (Detection Head): รวมคุณลักษณะเพื่อทำนายคลาสและกรอบตำแหน่ง",
    ])

    D.h("5. สถาปัตยกรรม YOLO และการตรวจจับแบบขั้นตอนเดียว", 2)
    D.p("YOLO เป็นสถาปัตยกรรมตรวจจับวัตถุที่ทำนายกรอบและคลาสในครั้งเดียว จึงเร็วและเหมาะกับงานตรวจสอบ "
        "โครงงานนี้ใช้ YOLO11 รุ่นเล็ก (yolo11n) ซึ่งมีพารามิเตอร์น้อย เหมาะกับเครื่องที่มีทรัพยากรจำกัด "
        "และได้ผลใกล้เคียงรุ่นใหญ่กว่าในโดเมนนี้ ตัวอย่างผลการตรวจจับของระบบแสดงในภาพ 2")
    D.figure("res_rust.jpg", "ตัวอย่างผลการตรวจจับตำหนิของระบบ — สนิม (rust) บนเหล็กเส้น "
             "พร้อมกรอบ ป้ายชนิดภาษาไทย และค่าความมั่นใจ", width=4.6)

    D.h("6. Material Segmentation (DMS46) สำหรับการระบุบริเวณโลหะ", 2)
    D.p("DMS46 (Dense Material Segmentation, Apple) เป็นแบบจำลองแบ่งส่วนภาพตามชนิดวัสดุ 46 ประเภทที่ฝึกมา "
        "แล้ว โครงงานใช้เฉพาะผลของคลาส “Metal” เพื่อระบุบริเวณวัสดุโลหะและตัดพื้นหลังที่ไม่ใช่เหล็กออกก่อน "
        "ส่งให้ขั้นตอนตรวจจับตำหนิ")

    D.h("7. ขั้นตอนการประมวลผลภาพดิจิทัล (Image Pre-processing)", 2)
    D.bullets([
        "การแปลงช่องสีและการปรับขนาด: ทำให้ภาพมีรูปแบบมาตรฐานตามที่แบบจำลองต้องการ",
        "การแปลงเป็นภาพระดับเทา (Grayscale): ใช้ในการฝึกขั้นที่ 2 เพื่อลดการเรียนรู้ทางลัดจากโทนสีที่แยก "
        "โดเมนของชุดข้อมูล (NEU เป็นภาพขาวดำ ส่วนสนิม/รอยแตกเป็นภาพสี)",
        "การทำ Normalize: ปรับค่าพิกเซลให้อยู่ในช่วงที่เหมาะสมก่อนเข้าสู่แบบจำลอง",
    ])

    D.h("8. การประเมินระดับความรุนแรงและการอธิบายผล (Severity & Explainability)", 2)
    D.p("การอธิบายผลในระบบปัจจุบันเป็นแบบเชิงภาพและเชิงตำแหน่ง ได้แก่ กรอบตำหนิ ค่าความมั่นใจ และระดับ "
        "ความเสี่ยงรายคลาส ส่วนการประเมิน Severity เชิงปริมาณ (คำนวณ Defect Area Ratio เทียบเกณฑ์ แล้วจัด "
        "เป็น Low/Medium/High และตัดสิน Pass/Fail) เป็นแนวทางที่ออกแบบไว้เป็นงานพัฒนาต่อ")

    D.h("9. เครื่องมือ ภาษา และไลบรารีที่เกี่ยวข้อง", 2)
    D.bullets([
        "ภาษา Python 3.11 เป็นภาษาหลักในการพัฒนา",
        "PyTorch 2.6 (CUDA 12.4) และ Ultralytics YOLO 8.4 สำหรับฝึกและรันแบบจำลอง",
        "OpenCV สำหรับการประมวลผลภาพ และ NumPy สำหรับการคำนวณเชิงเมทริกซ์",
        "Gradio สำหรับสร้างหน้าเว็บต้นแบบ และ matplotlib / python-docx สำหรับสร้างรูปและเอกสาร",
        "Roboflow สำหรับจัดการชุดข้อมูล การติดป้ายกำกับ และการทำ Data Augmentation",
        "Git และ GitHub สำหรับควบคุมเวอร์ชันและทำงานร่วมกัน",
    ])

    D.h("10. งานวิจัยที่เกี่ยวข้อง (Related Work)", 2)
    D.bullets([
        "Wang et al. (2023) ใช้ YOLOv5 ตรวจจับตำหนิบนแถบเหล็ก แสดงว่ากลุ่ม YOLO เร็วและแม่นยำเหมาะกับการ "
        "ตรวจสอบแบบเรียลไทม์",
        "Mao et al. (2025) ยืนยันประสิทธิภาพของ YOLO ในการตรวจจับความผิดปกติของชิ้นส่วนในสายการผลิต",
        "Luo et al. (2025) เสนอ DScanNet ที่ใช้ Selective State Space Models ตรวจจับตำหนิขนาดเล็กโดยใช้ "
        "ทรัพยากรน้อยลง สะท้อนแนวโน้มโมเดลที่เบาและเร็ว",
        "Mei et al. (2025) และ Huber et al. (2025) ศึกษาการใช้ข้อมูลสังเคราะห์และ Domain Randomization "
        "เพื่อแก้ปัญหา Class Imbalance และช่องว่างเชิงโดเมน",
        "Schmitt & Stief (2024) พัฒนาระบบ End-to-End Visual Quality Inspection ที่มี pipeline ตั้งแต่รับ "
        "ภาพจนถึงการตัดสิน Pass/Fail ในสภาพการผลิตจริง",
    ])

    D.h("11. การสังเคราะห์เชิงประเด็น (Thematic Synthesis)", 2)
    D.bullets([
        "ประเด็นที่ 1 — โมเดลที่เบาและเร็ว: งานวิจัยเปลี่ยนจาก CNN ขนาดใหญ่ไปสู่ตัวตรวจจับที่ประมวลผลไว "
        "(YOLO, SSM/Mamba)",
        "ประเด็นที่ 2 — ความน่าเชื่อถือและการอธิบายได้ (XAI): อุตสาหกรรมต้องการคำอธิบายเชิงภาพและเชิง "
        "ปริมาณ เช่น ตำหนิอยู่ที่ใด ขนาดเท่าใด และรุนแรงระดับใด",
        "ประเด็นที่ 3 — การนำไปใช้แบบครบวงจรและ Data Pipeline: ต้องมี pipeline ข้อมูลที่ดีและการบันทึก "
        "Metadata เพื่อเชื่อมต่อกับแดชบอร์ดและระบบแจ้งเตือน",
    ])

    D.h("12. ช่องว่างของงานวิจัย (Research Gap)", 2)
    D.bullets([
        "งานส่วนใหญ่หยุดที่การวาด Bounding Box แต่ไม่นำ Metadata มาคำนวณต่อเพื่ออธิบายระดับความรุนแรง "
        "อย่างเป็นรูปธรรม",
        "การเชื่อมจากการตรวจจับไปสู่ตรรกะตัดสิน Pass/Fail และการแจ้งเตือนที่ใช้งานได้จริงยังมีน้อย",
        "ช่องว่างเชิงโดเมน (Domain Gap) ระหว่างชุดข้อมูลแบบแล็บกับภาพถ่ายจริงยังเป็นอุปสรรคสำคัญที่ต้อง "
        "รายงานอย่างตรงไปตรงมา",
    ])

    D.h("13. ทฤษฎีเกี่ยวกับการประเมินความพึงพอใจ", 2)
    D.p("การประเมินความพึงพอใจใช้แบบสอบถามมาตราส่วนประมาณค่า 5 ระดับ (Likert Scale) และแปลผลค่าเฉลี่ยตาม "
        "เกณฑ์ในตาราง 2")
    D.table(["ช่วงค่าเฉลี่ย", "ระดับความพึงพอใจ"], [
        ["4.50 – 5.00", "มากที่สุด"], ["3.50 – 4.49", "มาก"], ["2.50 – 3.49", "ปานกลาง"],
        ["1.50 – 2.49", "น้อย"], ["1.00 – 1.49", "น้อยที่สุด"],
    ], "เกณฑ์การแปลผลค่าเฉลี่ยระดับความพึงพอใจของผู้ใช้งาน", col_w=[5, 6])

    D.h("14. การเปรียบเทียบระบบและเครื่องมือที่เกี่ยวข้อง", 2)
    D.table(["ระบบ/แพลตฟอร์ม", "หน้าที่หลัก", "จุดเด่น", "สิ่งที่โครงงานนำมาประยุกต์"], [
        ["Roboflow", "จัดการ Dataset และ Label", "ใช้งานง่าย มี Augmentation และ Dataset Versioning",
         "ใช้เตรียมและติดป้ายกำกับชุดข้อมูล"],
        ["YOLO-based Inspection", "ตรวจจับตำหนิจากภาพ", "เร็ว แม่นยำ เหมาะกับงานตรวจสอบ",
         "ใช้เป็นแบบจำลองตรวจจับตำหนิ (Stage 2)"],
        ["DMS46 (Material Seg.)", "แบ่งส่วนภาพตามชนิดวัสดุ", "แยกวัสดุโลหะจากพื้นหลังในภาพระดับสถานที่",
         "ใช้เป็นตัวระบุบริเวณโลหะ (Stage 1)"],
        ["Visual Quality Inspection", "ตรวจสอบคุณภาพอัตโนมัติแบบครบวงจร", "มี pipeline ตั้งแต่ภาพจนถึง Pass/Fail",
         "นำแนวคิด pipeline และการบันทึกผลมาต่อยอด (Future Work)"],
    ], "การเปรียบเทียบระบบและเครื่องมือที่เกี่ยวข้องกับโครงงาน", col_w=[3.4, 3.2, 4.2, 4.0])

    return D


if __name__ == "__main__":
    import thesis_content as TC
    print("1) สร้างไดอาแกรมขาวดำ ...")
    figs = {
        "framework": fig_framework(),
        "flowmain": fig_flow_main(),
        "flowpre": fig_flow_pre(),
        "classdist": fig_class_dist(),
    }
    figs.update(gen_reused_figs())
    print("2) ประกอบเล่ม ...")
    D = build_front_and_ch12(figs)
    TC.chapter3(D, figs)
    TC.bibliography(D)
    D.page_numbers()
    D.save()
    print("   ตาราง:", D.tab_n, " รูป:", D.fig_n)
    print("\nเปิดใน Word แล้วกด F9 (Update Field) ทั้งเอกสาร เพื่อสร้างสารบัญ/สารบัญตาราง/สารบัญรูปภาพ")
